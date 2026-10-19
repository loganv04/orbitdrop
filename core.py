"""Conversion engine, independent of the desktop interface.

Source files are read-only. Outputs are created exclusively, never overwritten.
PDFium calls share a lock because PDFium is not thread-safe.
"""

from dataclasses import dataclass
from pathlib import Path
import math
import threading
import warnings

from PIL import Image, ImageDraw, ImageFilter, ImageOps
import pypdfium2 as pdfium
from pypdf import PdfReader
from docx import Document

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
SUPPORTED_EXTENSIONS = IMAGE_EXTENSIONS | {".pdf"}
IMAGE_FORMATS = ("PNG", "JPG", "WEBP", "TIFF", "BMP")
MAX_PIXELS = 40_000_000
MAX_FILE_BYTES = 100 * 1024 * 1024
MAX_PAGES = 100
Image.MAX_IMAGE_PIXELS = MAX_PIXELS
PDF_LOCK = threading.RLock()


class ConversionError(ValueError):
    """A useful, user-facing conversion failure."""


@dataclass(frozen=True)
class Options:
    quality: int = 90
    dpi: int = 144

    def validate(self):
        if not 10 <= self.quality <= 100:
            raise ConversionError("Quality must be between 10 and 100.")
        if not 72 <= self.dpi <= 300:
            raise ConversionError("PDF resolution must be between 72 and 300 DPI.")


def validate_source(path):
    path = Path(path)
    if not path.is_file():
        raise ConversionError("The source file no longer exists.")
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ConversionError("Choose a PNG, JPG, WEBP, BMP, TIFF, or PDF file.")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ConversionError("Choose a file smaller than 100 MB.")
    return path


def formats_for(path):
    return IMAGE_FORMATS + (("TXT", "DOCX") if Path(path).suffix.lower() == ".pdf" else ("PDF",))


def check_dimensions(width, height):
    if width < 1 or height < 1 or width * height > MAX_PIXELS:
        raise ConversionError("This image exceeds the 40-megapixel limit. Try a lower PDF DPI.")


def load_image(path):
    """Read the first frame and honor EXIF orientation."""
    validate_source(path)
    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(path) as image:
            check_dimensions(*image.size)
            return ImageOps.exif_transpose(image).convert("RGBA")


def render_pdf_page(path, page_number=0, dpi=144):
    validate_source(path)
    with PDF_LOCK:
        try:
            with pdfium.PdfDocument(str(path)) as pdf:
                if not 0 <= page_number < len(pdf):
                    raise ConversionError("That PDF page does not exist.")
                page = pdf[page_number]
                try:
                    width, height = page.get_size()
                    check_dimensions(math.ceil(width * dpi / 72), math.ceil(height * dpi / 72))
                    bitmap = page.render(scale=dpi / 72)
                    try:
                        return bitmap.to_pil().convert("RGBA").copy()
                    finally:
                        bitmap.close()
                finally:
                    page.close()
        except pdfium.PdfiumError as exc:
            raise ConversionError("Cannot open this PDF. It may be damaged or password-protected.") from exc


def preview(path):
    image = render_pdf_page(path, dpi=90) if Path(path).suffix.lower() == ".pdf" else load_image(path)
    image.thumbnail((640, 480), Image.Resampling.LANCZOS)
    return image


def exclusive_output(directory, stem, extension):
    """Reserve a unique filename atomically, including collisions across processes."""
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    for number in range(10000):
        suffix = "" if number == 0 else f"_{number}"
        path = directory / f"{stem}{suffix}.{extension.lower()}"
        try:
            return path, path.open("xb")
        except FileExistsError:
            continue
    raise ConversionError("Too many files with the same name in this output folder.")


def flatten(image, color="white"):
    rgba = image.convert("RGBA")
    base = Image.new("RGB", rgba.size, color)
    base.paste(rgba, mask=rgba.getchannel("A"))
    return base


def save_image(image, directory, stem, target, options=Options()):
    options.validate()
    target = target.upper()
    if target not in IMAGE_FORMATS + ("PDF",):
        raise ConversionError("Unsupported image output format.")
    check_dimensions(*image.size)
    path, handle = exclusive_output(directory, stem, target)
    try:
        with handle:
            output = flatten(image) if target in {"JPG", "PDF", "BMP"} else image.convert("RGBA")
            # Remove inherited metadata; retain pixels only.
            clean = Image.new(output.mode, output.size)
            clean.paste(output)
            kwargs = {}
            if target in {"JPG", "WEBP"}:
                kwargs["quality"] = options.quality
            if target == "JPG":
                kwargs["optimize"] = True
            if target == "PNG":
                kwargs["optimize"] = True
            if target == "PDF":
                kwargs["resolution"] = options.dpi
            clean.save(handle, format="JPEG" if target == "JPG" else target, **kwargs)
        return path
    except Exception:
        path.unlink(missing_ok=True)
        raise


def extract_pdf_text(path):
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise ConversionError("Password-protected PDFs are not supported.")
    if len(reader.pages) > MAX_PAGES:
        raise ConversionError("Choose a PDF with 100 pages or fewer.")
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    if not any(pages):
        raise ConversionError("No selectable text found. Scanned PDFs need OCR, which is not included.")
    return pages


def convert(path, target, directory, options=Options(), progress=None):
    """Convert one file. A PDF produces one image per page; failures roll back this job."""
    options.validate()
    path = validate_source(path)
    target = target.upper()
    if target not in formats_for(path):
        raise ConversionError(f"{path.suffix.upper()} to {target} is not supported.")
    outputs = []
    try:
        if path.suffix.lower() != ".pdf":
            outputs.append(save_image(load_image(path), directory, path.stem, target, options))
        elif target in {"TXT", "DOCX"}:
            pages = extract_pdf_text(path)
            destination, handle = exclusive_output(directory, path.stem, target)
            outputs.append(destination)
            with handle:
                if target == "TXT":
                    handle.write("\n\n\f\n\n".join(pages).encode("utf-8"))
                else:
                    document = Document()
                    for index, text in enumerate(pages):
                        if index:
                            document.add_page_break()
                        for line in text.splitlines():
                            document.add_paragraph(line)
                    document.save(handle)
        else:
            with PDF_LOCK:
                with pdfium.PdfDocument(str(path)) as pdf:
                    count = len(pdf)
            if count == 0 or count > MAX_PAGES:
                raise ConversionError("Choose a PDF with 1–100 pages.")
            for index in range(count):
                if progress:
                    progress(f"Rendering page {index + 1} of {count}…")
                image = render_pdf_page(path, index, options.dpi)
                outputs.append(save_image(image, directory, f"{path.stem}_page_{index + 1:03}", target, options))
        return outputs
    except Exception:
        for output in outputs:
            output.unlink(missing_ok=True)
        raise


def crop_image(image, box):
    left, top, right, bottom = map(int, box)
    if not (0 <= left < right <= image.width and 0 <= top < bottom <= image.height):
        raise ConversionError("Crop selection must be inside the image and have a nonzero size.")
    return image.crop((left, top, right, bottom))


PALETTES = {
    "Sunset": ("#ffb35c", "#ea6787"),
    "Ocean": ("#8ae5da", "#5572d8"),
    "Lavender": ("#d9bcff", "#8b7bd7"),
    "Slate": ("#526379", "#1d2738"),
}


def add_frame(image, padding=80, radius=24, palette="Sunset"):
    if palette not in PALETTES or not 0 <= padding <= 400 or not 0 <= radius <= 200:
        raise ConversionError("Invalid frame settings.")
    width, height = image.width + padding * 2, image.height + padding * 2
    check_dimensions(width, height)
    from PIL import ImageColor
    start, end = [ImageColor.getrgb(color) for color in PALETTES[palette]]
    ramp = Image.new("RGB", (1, height))
    ramp.putdata([tuple(round(a + (b - a) * y / max(1, height - 1)) for a, b in zip(start, end)) for y in range(height)])
    background = ramp.resize((width, height)).convert("RGBA")
    mask = Image.new("L", image.size)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, image.width - 1, image.height - 1), radius=min(radius, min(image.size) // 2), fill=255)
    from PIL import ImageChops
    mask = ImageChops.multiply(mask, image.convert("RGBA").getchannel("A"))
    if padding:
        shadow_mask = Image.new("L", background.size)
        shadow_mask.paste(mask, (padding, padding + min(12, padding // 3)))
        shadow_mask = shadow_mask.filter(ImageFilter.GaussianBlur(min(24, padding // 3)))
        shadow = Image.new("RGBA", background.size, (25, 17, 40, 0))
        shadow.putalpha(shadow_mask.point(lambda p: round(p * 0.28)))
        background = Image.alpha_composite(background, shadow)
    background.paste(image.convert("RGBA"), (padding, padding), mask)
    return background
